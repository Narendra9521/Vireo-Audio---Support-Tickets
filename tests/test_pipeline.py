"""Run: python tests/test_pipeline.py   (no pytest needed; pytest also works)"""
import os, sys, tempfile
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
import pandas as pd
from openpyxl import load_workbook
from vireo_refunds.clean import clean_tickets
from vireo_refunds.classify import classify_text, notes_mention_replacement, replacement_evidence
from vireo_refunds import __main__ as cli


def mk(rows):
    base = dict(created_at="2025-02-01 10:00", first_response_at=None, resolved_at=None, status="resolved",
                refund_reason_code=None, replacement_issued="N", customer_message="", agent_notes="")
    df = pd.DataFrame([{**base, **r} for r in rows])
    df["created_at"] = pd.to_datetime(df["created_at"])
    return df


# --- amount normalisation + duplicates -------------------------------------------------------
def test_legacy_divided_and_deduped():
    t = mk([dict(ticket_id="T1", source_system="helpdesk", refund_amount_inr=900.0),
            dict(ticket_id="T1", source_system="legacy_fd", refund_amount_inr=90000.0),
            dict(ticket_id="T2", source_system="legacy_fd", refund_amount_inr=50000.0)])
    clean, bridge = clean_tickets(t)
    assert len(clean) == 2
    assert clean["refund_inr"].sum() == 900 + 500
    assert bridge[0][1] == 140900 and bridge[-1][1] == 1400      # raw -> clean bridge


def test_helpdesk_copy_kept_regardless_of_row_order():
    t = mk([dict(ticket_id="T1", source_system="legacy_fd", refund_amount_inr=90000.0),
            dict(ticket_id="T1", source_system="helpdesk", refund_amount_inr=900.0)])
    clean, _ = clean_tickets(t)
    assert clean.iloc[0]["source_system"] == "helpdesk" and clean.iloc[0]["refund_inr"] == 900


def test_blank_refund_stays_blank():
    t = mk([dict(ticket_id="T1", source_system="helpdesk", refund_amount_inr=None)])
    clean, _ = clean_tickets(t)
    assert pd.isna(clean.iloc[0]["refund_inr"])


# --- month assignment / timezone ---------------------------------------------------------------
def test_month_uses_created_at_as_exported_ist_no_shift():
    # Decision (see README): created_at is already IST in both systems, so no conversion is applied.
    t = mk([dict(ticket_id="T1", source_system="helpdesk", refund_amount_inr=100.0, created_at="2025-01-31 23:30"),
            dict(ticket_id="T2", source_system="helpdesk", refund_amount_inr=100.0, created_at="2025-02-01 00:10")])
    clean, _ = clean_tickets(t)
    assert clean.set_index("ticket_id").loc["T1", "month"] == "2025-01"
    assert clean.set_index("ticket_id").loc["T2", "month"] == "2025-02"


# --- reason classification -----------------------------------------------------------------------
def test_classifier_examples():
    assert classify_text("charged twice for one order", "")[0] == "DUP-PAYMENT"
    assert classify_text("cancel button greyed out", "")[0] == "CANCEL"
    assert classify_text("coupon not working", "refunded")[0] == "PRICE-ADJ"
    assert classify_text("order not delivered", "")[0] == "LOST-TRANSIT"
    assert classify_text("bluetooth pairing fails", "")[0] == "FAULT-DOA-WTY"
    assert classify_text("where is my money for the return", "")[0] == "RETURN-QC-OK"
    assert classify_text("", "")[0] == "UNCLASSIFIED"
    assert classify_text(None, float("nan"))[0] == "UNCLASSIFIED"


def test_agent_note_beats_message():
    assert classify_text("something odd", "double charge, refunded")[0] == "DUP-PAYMENT"


# --- refund + replacement detection --------------------------------------------------------------
def test_replacement_positive():
    for n in ["refund + rplc, cx escalation avoided",
              "Both refund and replacement given as cx threatened social media.",
              "Refunded Rs 1499 & also raised replacement, cx was very upset.",
              "Cx got the amount back, also sending a new set to keep him happy.",
              "Reversed the payment and dispatched a new one under RMA."]:
        assert notes_mention_replacement(n), n


def test_replacement_negative():
    for n in ["Cx asked for rplc + rfnd, explained policy, rfnd only.",
              "rplc was offered earlier by chat team, cx opted for rfnd.",
              "Refund issued; replacement not applicable (out of stock).",
              "refunded. rplc request rejected as per policy.",
              "Offered replacement, cx preferred refund.", "", None, float("nan")]:
        assert not notes_mention_replacement(n), n


def test_replacement_evidence_tiers():
    assert replacement_evidence("refund + rplc", True) == "flag + notes"
    assert replacement_evidence("sorted", True) == "flag only"
    assert replacement_evidence("refund + rplc, cx escalation avoided", False) == "notes, explicit refund+replacement"
    assert replacement_evidence("Replacement dispatched.", False) == "notes, replacement only"
    assert replacement_evidence("sorted", False) == ""


# --- aggregation / reconciliation, end to end on a tiny synthetic data set -----------------------
def test_end_to_end_totals_reconcile():
    with tempfile.TemporaryDirectory() as d:
        data, out = os.path.join(d, "data"), os.path.join(d, "out")
        os.makedirs(data)
        cols = "ticket_id,created_at,first_response_at,resolved_at,status,channel,customer_id,order_id,product_sku,category,priority,assigned_team,agent_id,transfers,csat_score,refund_amount_inr,refund_reason_code,replacement_issued,customer_message,agent_notes,source_system"
        rows = [
            "T1,2025-01-10 10:00,,,resolved,chat,C1,O1,VA-EB-PL2,Other,Low,Chat Frontline,A1,0,,2000,GW-OTHER,N,charged twice,refunded,helpdesk",
            "T1,2025-01-10 10:00,,,resolved,chat,C1,O1,VA-EB-PL2,Other,Low,Chat Frontline,A1,0,,200000,GW-OTHER,N,charged twice,refunded,legacy_fd",
            "T2,2025-02-03 11:00,,,resolved,chat,C2,O2,VA-EB-PL2,Other,Low,Chat Frontline,A1,0,,3000,CANCEL,Y,cancel order,refund + rplc,helpdesk",
            "T3,2025-02-04 11:00,,,resolved,chat,C3,,VA-EB-PL2,Other,Low,Chat Frontline,A1,0,,,,N,hello,ok,helpdesk",
        ]
        open(f"{data}/tickets.csv", "w").write(cols + "\n" + "\n".join(rows) + "\n")
        open(f"{data}/agents.csv", "w").write("agent_id,name,site,team,shift,tier,from_date,to_date\nA1,Test Agent,Indore,Chat Frontline,Day,1,2024-01-01,\n")
        open(f"{data}/orders.csv", "w").write("order_id,customer_id,sku,order_date,channel,qty,order_value_inr,lot_code\nO1,C1,VA-EB-PL2,2025-01-01,vireo.in,1,3499,L1\nO2,C2,VA-EB-PL2,2025-01-02,vireo.in,1,3499,L1\n")
        open(f"{data}/products.csv", "w").write("sku,product_name,family,launch_date,unit_cost_inr,retail_price_inr,warranty_months\nVA-EB-PL2,Pulse 2,earbuds,2025-07-15,1480,3499,12\n")
        cli.main(["--data", data, "--out", out])
        wb = load_workbook(f"{out}/refund_summary.xlsx")
        ws = wb["By reason code (agent)"]
        body = [c.value for row in ws.iter_rows(min_row=2, max_row=ws.max_row - 1, min_col=2, max_col=ws.max_column - 1) for c in row]
        assert sum(v for v in body if isinstance(v, (int, float))) == 2000 + 3000
        rec = {r[0].value: r[1].value for r in wb["Reconciliation"].iter_rows(min_row=2)}
        assert rec["Clean refund total"] == 5000
        t2 = pd.read_csv(f"{out}/tickets_clean.csv").set_index("ticket_id")
        assert bool(t2.loc["T2", "refund_and_replacement"]) and t2.loc["T2", "replacement_cost_inr"] == 1480 + 340
        assert not bool(t2.loc["T1", "refund_and_replacement"])
        assert t2.loc["T1", "inferred_reason"] == "DUP-PAYMENT"


def test_real_data_timezone_evidence():
    """Skipped when the client CSV is absent. Voice callbacks are only taken 08:00-22:00 IST (policy s2); if
    created_at were UTC the hours would sit at ~02:30-16:30. Guards the 'created_at is IST, no shift' decision."""
    path = os.path.join(os.path.dirname(__file__), "..", "data", "tickets.csv")
    if not os.path.exists(path):
        return
    v = pd.read_csv(path, usecols=["channel", "created_at"], parse_dates=["created_at"])
    h = v[v["channel"] == "voice"]["created_at"].dt.hour
    assert h.min() >= 8 and h.max() < 22, (h.min(), h.max())


if __name__ == "__main__":
    n = 0
    for name, fn in list(globals().items()):
        if name.startswith("test_"):
            fn(); n += 1; print("ok", name)
    print(f"{n} tests passed")
